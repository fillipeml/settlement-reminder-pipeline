"""The client directory (a spreadsheet exported from the case system)."""

from __future__ import annotations

from openpyxl import Workbook

from settlement_reminders.ingest import ExcelClientDirectory

HEADERS = ["Name", "Group", "Tax ID", "Entity type", "E-mails"]


def _make_directory(tmp_path, rows):
    wb = Workbook()
    ws = wb.active
    ws.append(HEADERS)
    for row in rows:
        ws.append(row)
    path = tmp_path / "clients.xlsx"
    wb.save(path)
    return path


def test_reads_contacts(tmp_path):
    path = _make_directory(
        tmp_path,
        [
            (
                "GRUPO MODELO CONSTRUTORA E INCORPORADORA LTDA",
                "Grupo Modelo",
                "",
                "Legal entity",
                "finance@example.com; legal@example.com",
            ),
            ("CONSTRUTORA SEM EMAIL LTDA", None, "", "Legal entity", None),
        ],
    )

    contacts = ExcelClientDirectory(path).fetch_contacts()

    assert len(contacts) == 2
    first = contacts[0]
    assert first.group == "Grupo Modelo"
    assert [str(e) for e in first.emails] == ["finance@example.com", "legal@example.com"]
    assert contacts[1].emails == []


def test_portuguese_headers_are_accepted(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.append(["Nome", "Nome > Grupo Empresarial", "CPF/CNPJ", "Tipo Pessoa", "E-mails"])
    ws.append(["CLIENTE X", "G", "", "Jurídica", "a@example.com"])
    path = tmp_path / "relatorio.xlsx"
    wb.save(path)
    contacts = ExcelClientDirectory(path).fetch_contacts()
    assert contacts[0].group == "G"
    assert [str(e) for e in contacts[0].emails] == ["a@example.com"]


def test_invalid_email_is_discarded_without_dropping_the_row(tmp_path):
    path = _make_directory(tmp_path, [("CLIENT X", None, None, None, "broken@; valid@example.com")])

    contacts = ExcelClientDirectory(path).fetch_contacts()

    assert [str(e) for e in contacts[0].emails] == ["valid@example.com"]


def test_row_without_a_name_is_ignored(tmp_path):
    path = _make_directory(
        tmp_path,
        [(None, None, None, None, "orphan@example.com"), ("CLIENT Y", None, None, None, None)],
    )

    contacts = ExcelClientDirectory(path).fetch_contacts()

    assert [c.name for c in contacts] == ["CLIENT Y"]


def test_missing_required_header_fails_with_a_clear_message(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.append(["Name", "Tax ID"])  # no e-mail column
    path = tmp_path / "invalid.xlsx"
    wb.save(path)

    try:
        ExcelClientDirectory(path).fetch_contacts()
    except ValueError as exc:
        assert "emails" in str(exc)
    else:
        raise AssertionError("should fail without the e-mail column")
