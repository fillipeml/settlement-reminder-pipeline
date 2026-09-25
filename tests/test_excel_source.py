from datetime import date
from decimal import Decimal

import pytest
from openpyxl import Workbook

from settlement_reminders.sources.excel import ExcelSource
from tests.factories import CASE_B

COLUMNS = [
    "Case number",
    "Claimant",
    "Payer",
    "Beneficiary",
    "Honorific",
    "Client emails",
    "Lawyer email",
    "Installment amount",
    "Installment count",
    "First due",
    "Payment method",
    "Payment details",
    "Late clause",
    "Active",
]


def _make(tmp_path, rows):
    wb = Workbook()
    ws = wb.active
    ws.append(COLUMNS)
    for r in rows:
        ws.append(r)
    p = tmp_path / "agreements.xlsx"
    wb.save(p)
    return p


def _row(**o):
    base = dict(
        case=CASE_B,
        claimant="Sicrana",
        payer="Empresa",
        beneficiary="Sicrana",
        honorific="Sra.",
        emails="a@x.com, b@y.com",
        lawyer="lawyer@lawfirm.example",
        amount="2166,66",
        count=3,
        first=date(2026, 6, 15),
        method="depósito ou transferência",
        details="Banco Exemplo\nPIX: 000",
        clause="em caso de atraso, multa",
        active="Yes",
    )
    base.update(o)
    return [
        base[k]
        for k in (
            "case",
            "claimant",
            "payer",
            "beneficiary",
            "honorific",
            "emails",
            "lawyer",
            "amount",
            "count",
            "first",
            "method",
            "details",
            "clause",
            "active",
        )
    ]


def test_reads_an_agreement_and_expands_it(tmp_path):
    path = _make(tmp_path, [_row()])
    agreements = list(ExcelSource(str(path)).fetch_agreements())
    assert len(agreements) == 1
    a = agreements[0]
    assert a.installment_amount == Decimal("2166.66")
    assert [str(e) for e in a.client_emails] == ["a@x.com", "b@y.com"]
    assert len(a.installments()) == 3


def test_portuguese_headers_are_accepted(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.append(
        [
            "Processo",
            "Autor",
            "Reu",
            "Beneficiario",
            "Tratamento",
            "Emails Cliente",
            "Email Advogado",
            "Valor Parcela",
            "Num Parcelas",
            "Primeira Parcela",
            "Forma Pagamento",
            "Dados Pagamento",
            "Clausula Mora",
            "Ativo",
        ]
    )
    ws.append(_row(active="Sim"))
    p = tmp_path / "acordos.xlsx"
    wb.save(p)
    assert len(list(ExcelSource(str(p)).fetch_agreements())) == 1


def test_ignores_inactive_rows(tmp_path):
    path = _make(tmp_path, [_row(active="No")])
    assert list(ExcelSource(str(path)).fetch_agreements()) == []


def test_missing_required_column(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.append(["Case number", "Claimant"])
    ws.append(["123", "Someone"])
    p = tmp_path / "bad.xlsx"
    wb.save(p)
    with pytest.raises(ValueError):
        list(ExcelSource(str(p)).fetch_agreements())
