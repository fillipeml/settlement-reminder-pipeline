"""Data source: an Excel spreadsheet (.xlsx) mirroring the SharePoint list.

Useful for tests and small teams. The columns follow the same schema as the SharePoint
list (one agreement per row); headers are matched by tolerant aliases in English and
Portuguese.
"""

from __future__ import annotations

import logging
import unicodedata
from collections.abc import Iterable
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook
from pydantic import ValidationError

from settlement_reminders.models import Agreement
from settlement_reminders.sources.base import DataSource

logger = logging.getLogger(__name__)

# logical field -> possible headers (normalised: lower case, no accents)
COLUMN_ALIASES: dict[str, tuple[str, ...]] = {
    "case_number": (
        "case number",
        "case",
        "processo",
        "numero processo",
        "n processo",
        "cnj",
        "numero do processo",
    ),
    "claimant": ("claimant", "autor", "reclamante", "exequente"),
    "payer": ("payer", "client", "reu", "reclamada", "cliente", "executado", "parte re"),
    "beneficiary_name": (
        "beneficiary",
        "beneficiary name",
        "beneficiario",
        "beneficiario nome",
        "nome beneficiario",
        "favorecido",
    ),
    "honorific": ("honorific", "tratamento"),
    "client_emails": (
        "client emails",
        "client email",
        "emails cliente",
        "email cliente",
        "emails",
        "e mails",
        "e mail cliente",
    ),
    "lawyer_email": (
        "lawyer email",
        "email advogado",
        "advogado email",
        "advogado responsavel",
        "email adv",
    ),
    "installment_amount": (
        "installment amount",
        "amount",
        "valor parcela",
        "valor da parcela",
        "valor",
    ),
    "installment_count": (
        "installment count",
        "installments",
        "num parcelas",
        "numero de parcelas",
        "quantidade parcelas",
        "qtd parcelas",
        "parcelas",
    ),
    "first_due": (
        "first due",
        "first due date",
        "primeira parcela",
        "data primeira parcela",
        "data 1a parcela",
        "vencimento primeira parcela",
        "1a parcela",
    ),
    "payment_method": ("payment method", "forma pagamento", "forma de pagamento"),
    "payment_details": (
        "payment details",
        "bank details",
        "dados pagamento",
        "dados bancarios",
        "dados de pagamento",
        "conta",
    ),
    "late_clause": (
        "late clause",
        "late payment clause",
        "clausula mora",
        "clausula de mora",
        "mora",
        "multa",
    ),
    "active": ("active", "ativo"),
}

REQUIRED = (
    "case_number",
    "claimant",
    "payer",
    "beneficiary_name",
    "client_emails",
    "installment_amount",
    "installment_count",
    "first_due",
    "payment_details",
    "late_clause",
)


def normalize_header(text: object) -> str:
    if text is None:
        return ""
    s = unicodedata.normalize("NFKD", str(text))
    s = "".join(c for c in s if not unicodedata.combining(c))
    s = s.lower().strip()
    for ch in ("_", "-", ".", "/"):
        s = s.replace(ch, " ")
    return " ".join(s.split())


def _parse_date(value: object) -> date | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    try:
        return datetime.fromisoformat(text.replace("Z", "+00:00")).date()
    except ValueError:
        pass
    for fmt in ("%d/%m/%Y", "%d-%m-%Y", "%Y-%m-%d", "%d/%m/%y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def _parse_decimal(value: object) -> Decimal | None:
    if value is None or value == "":
        return None
    if isinstance(value, int | float | Decimal):
        try:
            return Decimal(str(value))
        except InvalidOperation:
            return None
    text = str(value).strip()
    for token in ("R$", " ", "\xa0"):
        text = text.replace(token, "")
    text = text.replace(".", "").replace(",", ".")
    try:
        return Decimal(text)
    except InvalidOperation:
        return None


def _parse_bool(value: object) -> bool:
    if value is None:
        return True
    return normalize_header(value) in (
        "yes",
        "y",
        "sim",
        "s",
        "true",
        "1",
        "verdadeiro",
        "ativo",
        "active",
    )


def row_to_agreement(get) -> Agreement:
    """Builds an Agreement from a `get(field)` accessor."""
    due = _parse_date(get("first_due"))
    if due is None:
        raise ValueError("first_due missing or invalid")
    amount = _parse_decimal(get("installment_amount"))
    if amount is None:
        raise ValueError("installment_amount missing or invalid")
    count = get("installment_count")

    return Agreement(
        case_number=get("case_number"),
        claimant=get("claimant"),
        payer=get("payer"),
        beneficiary_name=get("beneficiary_name"),
        honorific=get("honorific") or "Sr.",
        client_emails=get("client_emails") or "",
        lawyer_email=(get("lawyer_email") or None),
        installment_amount=amount,
        installment_count=int(count),
        first_due=due,
        payment_method=get("payment_method") or "depósito ou transferência",
        payment_details=get("payment_details") or "",
        late_clause=get("late_clause") or "",
        active=_parse_bool(get("active")),
    )


def map_columns(header: Iterable[str], *, where: str = "spreadsheet") -> dict[str, str]:
    """Maps logical fields to the actual header names; raises when a required one is missing."""
    normalized = {normalize_header(h): h for h in header if h is not None}
    mapping: dict[str, str] = {}
    for field, aliases in COLUMN_ALIASES.items():
        for alias in aliases:
            if alias in normalized:
                mapping[field] = normalized[alias]
                break
    missing = [c for c in REQUIRED if c not in mapping]
    if missing:
        raise ValueError(
            f"Required columns not found in the {where}: {missing}. Headers read: {list(normalized)}"
        )
    return mapping


class ExcelSource(DataSource):
    """Reads agreements from an Excel spreadsheet (a mirror of the SharePoint list)."""

    name = "excel"

    def __init__(self, path: str) -> None:
        self.path = Path(path)

    def fetch_agreements(self) -> Iterable[Agreement]:
        if not self.path.exists():
            raise FileNotFoundError(f"Spreadsheet not found: {self.path}")

        wb = load_workbook(self.path, read_only=True, data_only=True)
        ws = wb.active
        rows = ws.iter_rows(values_only=True)
        try:
            header = next(rows)
        except StopIteration:
            logger.warning("Spreadsheet %s is empty.", self.path)
            return []

        by_name = map_columns([h for h in header if h is not None])
        index = {h: i for i, h in enumerate(header)}
        col = {field: index[name] for field, name in by_name.items()}
        agreements: list[Agreement] = []
        for line_no, row in enumerate(rows, start=2):
            if row is None or all(c is None or c == "" for c in row):
                continue

            def get(field: str, _row=row):
                idx = col.get(field)
                return _row[idx] if idx is not None and idx < len(_row) else None

            try:
                agreement = row_to_agreement(get)
            except (ValidationError, ValueError) as exc:
                logger.warning("Row %d ignored (%s): %s", line_no, self.path, exc)
                continue
            if agreement.active:
                agreements.append(agreement)

        wb.close()
        logger.info("Excel: %d active agreement(s) read from %s", len(agreements), self.path)
        return agreements
