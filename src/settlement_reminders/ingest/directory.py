"""Client directory (name -> e-mails), with pluggable sources.

Initial source: the client report exported from the case-management system as a
spreadsheet (columns: Name | Group | Tax ID | Entity type | E-mails). When the system's API
is connected, implement another `ClientDirectory`; the matcher does not change.
"""

from __future__ import annotations

import logging
from abc import ABC, abstractmethod
from pathlib import Path

from openpyxl import load_workbook
from pydantic import BaseModel, EmailStr, Field, TypeAdapter, ValidationError

from settlement_reminders.sources.excel import normalize_header

logger = logging.getLogger(__name__)

_EMAIL = TypeAdapter(EmailStr)

# logical field -> possible headers (normalised: lower case, no accents)
_COLUMNS: dict[str, tuple[str, ...]] = {
    "name": ("name", "client", "client name", "nome", "cliente", "nome cliente"),
    "group": (
        "group",
        "corporate group",
        "name > corporate group",
        "nome > grupo empresarial",
        "grupo empresarial",
        "grupo",
    ),
    "tax_id": ("tax id", "cpf cnpj", "cnpj", "cpf", "documento", "document"),
    "entity_type": ("entity type", "tipo pessoa"),
    "emails": ("emails", "e mails", "e mail", "email"),
}
_REQUIRED = ("name", "emails")


class ClientContact(BaseModel):
    """A client of the directory, with their contact e-mails."""

    name: str
    group: str = ""
    tax_id: str = ""
    entity_type: str = ""
    emails: list[EmailStr] = Field(default_factory=list)


class ClientDirectory(ABC):
    """Contract of a source of client contacts."""

    name: str = "unknown"

    @abstractmethod
    def fetch_contacts(self) -> list[ClientContact]:
        raise NotImplementedError


def _parse_emails(value: object) -> list[str]:
    """'a@x.com; b@y.com' -> ['a@x.com', 'b@y.com']; invalid ones are discarded."""
    if value is None:
        return []
    valid: list[str] = []
    for raw in str(value).replace(";", ",").split(","):
        address = raw.strip()
        if not address:
            continue
        try:
            valid.append(str(_EMAIL.validate_python(address)))
        except ValidationError:
            logger.warning("Invalid e-mail discarded from the directory: %r", address)
    return valid


class ExcelClientDirectory(ClientDirectory):
    """Reads the client report exported as a spreadsheet."""

    name = "excel"

    def __init__(self, path: str | Path) -> None:
        self.path = Path(path)

    def fetch_contacts(self) -> list[ClientContact]:
        if not self.path.exists():
            raise FileNotFoundError(f"Client directory not found: {self.path}")

        wb = load_workbook(self.path, read_only=True, data_only=True)
        ws = wb.active
        rows = ws.iter_rows(values_only=True)
        try:
            header = next(rows)
        except StopIteration:
            logger.warning("Directory %s is empty.", self.path)
            return []

        col = self._map_columns(header)
        contacts: list[ClientContact] = []
        for row in rows:
            if row is None:
                continue

            def get(field: str, _row=row):
                idx = col.get(field)
                return _row[idx] if idx is not None and idx < len(_row) else None

            name = get("name")
            if name is None or not str(name).strip():
                continue
            contacts.append(
                ClientContact(
                    name=str(name).strip(),
                    group=str(get("group") or "").strip(),
                    tax_id=str(get("tax_id") or "").strip(),
                    entity_type=str(get("entity_type") or "").strip(),
                    emails=_parse_emails(get("emails")),
                )
            )
        wb.close()

        without_email = sum(1 for c in contacts if not c.emails)
        logger.info(
            "Directory: %d client(s) read from %s (%d without e-mail)",
            len(contacts),
            self.path,
            without_email,
        )
        return contacts

    def _map_columns(self, header: tuple) -> dict[str, int]:
        normalized = {normalize_header(h): i for i, h in enumerate(header) if h is not None}
        mapping: dict[str, int] = {}
        for field, aliases in _COLUMNS.items():
            for alias in aliases:
                if alias in normalized:
                    mapping[field] = normalized[alias]
                    break
        missing = [c for c in _REQUIRED if c not in mapping]
        if missing:
            raise ValueError(
                f"Required columns not found in the directory: {missing}. Headers read: {list(normalized)}"
            )
        return mapping
